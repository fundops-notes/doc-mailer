"""v3.8.3: entry-point ownership, Excel tag roundtrip and persistent pool merges."""
from pathlib import Path
import pytest
import openpyxl

from test_v3_pool_gui import (app, no_messagebox, _pool_rec, _event, _walk,
                              _click_dialog_button, _pick_event_stub)
import lp_tool
import lp_v3


def _buttons(parent):
    return {str(w.cget("text")): w for w in _walk(parent)
            if w.winfo_class() == "TButton"}


def _read(path):
    book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        return [list(row) for row in book.active.iter_rows(values_only=True)]
    finally:
        book.close()


def _file(path, data):
    book = openpyxl.Workbook()
    for row in data:
        book.active.append(row)
    book.save(path)
    book.close()


def _import(app, monkeypatch, path):
    monkeypatch.setattr(lp_tool.filedialog, "askopenfilename", lambda **kw: str(path))

    def approve():
        dialog = next(w for w in app.root.winfo_children()
                      if w.winfo_class() == "Toplevel" and w.title() == "导入结果预览")
        button = next(w for name, w in _buttons(dialog).items() if name.startswith("导入无错误的"))
        button.invoke()

    app.root.after(80, approve)
    _buttons(app.home_tab_pool)["导入联系人 Excel/CSV"].invoke()


def test_entry_points_have_clear_ownership(app):
    home = _buttons(app.home_tab_ev)
    assert not {"导入收件人…", "导出收件人…", "导出全部收件人…"} & set(home)
    assert {"导出事件…", "导入事件备份…", "备份全部…"} <= set(home)
    menu = app._events_export_menu
    assert menu.entrycget(0, "label") == "导出各事件收件人明细…"
    assert not hasattr(app, "_lib_import_recipients")
    assert not hasattr(app, "_lib_export_recipients")
    assert {"导入联系人 Excel/CSV", "导出联系人 Excel", "下载联系人模板"} <= set(_buttons(app.home_tab_pool))
    app.open_event(_event(app, "目标事件")["id"])
    assert {"导入 Excel/CSV", "导出本事件收件人", "从收件人库添加…"} <= set(_buttons(app._wb_content))


def test_cross_event_export_uses_global_field_headers_and_keeps_event_rows():
    events = [{"name": "A", "recipients": [{"name": "同一人", "fields": {"Call_Amount": "100", "Due_Date": "2026/11/01"}}]},
              {"name": "B", "recipients": [{"name": "同一人", "fields": {"NAV": "1.2345", "Call_Amount": "200"}}]}]
    rows = lp_v3.events_recipients_to_rows(events)
    header = rows[0]
    assert len(rows) == 3 and all(len(row) == len(header) for row in rows)
    assert all(header) and len(set(header)) == len(header)
    a, b = (dict(zip(header, row)) for row in rows[1:])
    assert a["所属事件"] == "A" and b["所属事件"] == "B"
    assert a["Call_Amount"] == "100" and b["Call_Amount"] == "200"
    assert a["Due_Date"] == "2026/11/01" and b["Due_Date"] == ""
    assert a["NAV"] == "" and b["NAV"] == "1.2345"


def test_more_export_menu_writes_real_excel_with_named_columns(app, tmp_path, monkeypatch):
    _event(app, "事件A", [{"name": "同名", "fields": {"金额": "100"}}])
    _event(app, "事件B", [{"name": "同名", "fields": {"截止日期": "2026/11/01"}}])
    path = tmp_path / "events.xlsx"
    monkeypatch.setattr(lp_tool.filedialog, "asksaveasfilename", lambda **kw: str(path))
    app._events_export_menu.invoke("导出各事件收件人明细…")
    rows = _read(path)
    assert None not in rows[0]
    mapped = {r[0]: dict(zip(rows[0], r)) for r in rows[1:]}
    assert mapped["事件A"]["金额"] == "100"
    assert mapped["事件B"]["截止日期"] == "2026/11/01"


def test_cross_event_field_names_do_not_collide_with_fixed_headers():
    rows = lp_v3.events_recipients_to_rows([
        {"name": "A", "recipients": [{"name": "甲", "fields": {
            "附件路径": "业务字段", "所属事件": "自定义", "本事件字段：附件路径": "另一个"}}]}])
    assert len(set(rows[0])) == len(rows[0])
    values = dict(zip(rows[0], rows[1]))
    assert values["所属事件"] == "A"
    assert values["本事件字段：附件路径"] == "业务字段"
    assert values["本事件字段：所属事件"] == "自定义"
    assert "另一个" in values.values()


@pytest.mark.parametrize("text,expected", [
    ("华东；重点；华东", ["华东", "重点"]),
    ("华东,重点\n境外", ["华东", "重点", "境外"]),
    ("华东、重点;境外，长期", ["华东", "重点", "境外", "长期"]),
    ('["A,B", "[特殊]", "A,B"]', ["A,B", "[特殊]"]),
    ("", []),
])
def test_pool_tag_cell_parsing(text, expected):
    assert lp_v3.parse_pool_tags(text) == expected


def test_pool_validation_recognizes_tag_aliases_and_keeps_event_semantics():
    rows = [["收件人名称", "收件人", "TAGS", "联系人标签", "Call_Amount"],
            ["甲", "a@example.com", "重点;华东", "重点；境外", "100"]]
    pool = lp_v3.validate_pool_recipient_rows(rows)
    assert pool["items"][0]["record"]["tags"] == ["重点", "华东", "境外"]
    assert pool["items"][0]["record"]["fields"] == {}
    assert pool["ignored_cols"] == ["Call_Amount"]
    assert not pool["unknown_cols"] and not pool["field_cols"]
    assert any("Call_Amount" in note for note in pool["notes"])
    event = lp_v3.validate_recipient_rows(rows)
    assert event["items"][0]["record"]["fields"]["TAGS"] == "重点;华东"


def test_pool_roundtrip_preserves_special_tags_and_excludes_legacy_event_fields():
    rec = {"name": "甲", "to": "a@example.com", "open_pw": "open12", "owner_pw": "owner12",
           "tags": ["重点", "A,B", "[特殊]", "含；分号"], "fields": {"Call_Amount": "100"}}
    rows = lp_v3.pool_recipients_to_rows([rec])
    assert "标签" in rows[0] and "Call_Amount" not in rows[0]
    parsed = lp_v3.validate_pool_recipient_rows(rows)["items"][0]["record"]
    assert parsed["tags"] == rec["tags"] and parsed["owner_pw"] == "owner12"
    masked = lp_v3.pool_recipients_to_rows([rec], include_passwords=False)
    assert masked[1][rows[0].index("打开口令")] == ""
    assert masked[1][rows[0].index("权限口令")] == ""


def test_gui_pool_export_and_reimport_preserves_tags(app, tmp_path, monkeypatch):
    _pool_rec(app, "甲", tags=["重点", "华东", "A,B"], fields={"Old_Field": "100"})
    path = tmp_path / "pool.xlsx"
    monkeypatch.setattr(lp_tool.filedialog, "asksaveasfilename", lambda **kw: str(path))
    _buttons(app.home_tab_pool)["导出联系人 Excel"].invoke()
    assert "标签" in _read(path)[0] and "Old_Field" not in _read(path)[0]
    app.store = lp_v3.Store(tmp_path / "fresh_store")
    _import(app, monkeypatch, path)
    saved = app.store.list_pool()
    assert len(saved) == 1 and saved[0]["tags"] == ["重点", "华东", "A,B"]
    assert saved[0]["fields"] == {}


def test_gui_same_name_import_persists_cc_tags_and_keeps_passwords(app, tmp_path, monkeypatch):
    _pool_rec(app, "甲", to="existing@example.com", tags=["原标签"], open_pw="old-open", owner_pw="old-owner")
    path = tmp_path / "update.xlsx"
    _file(path, [["收件人名称", "收件人", "抄送", "打开口令", "权限口令", "标签", "Call_Amount"],
                 ["甲", "new@example.com", "cc@example.com", "new-open", "new-owner", "新增", "100"]])
    _import(app, monkeypatch, path)
    saved = app.store.list_pool()[0]
    assert saved["cc"] == "cc@example.com" and saved["to"] == "existing@example.com"
    assert saved["tags"] == ["原标签", "新增"] and saved["fields"] == {}
    assert saved["open_pw"] == "old-open" and saved["owner_pw"] == "old-owner"
    path2 = tmp_path / "blank_tag.xlsx"
    _file(path2, [["收件人名称", "收件人", "标签"], ["甲", "a@example.com", ""]])
    _import(app, monkeypatch, path2)
    assert app.store.list_pool()[0]["tags"] == ["原标签", "新增"]


@pytest.mark.parametrize("entry", ["_pool_from_event", "_wb_save_to_pool"])
def test_event_to_pool_entry_points_also_persist_same_name_updates(app, entry, monkeypatch):
    _pool_rec(app, "甲", open_pw="old-open", owner_pw="old-owner")
    ev = _event(app, "来源事件", [{"name": "甲", "to": "new@example.com", "cc": "cc@example.com",
                                  "open_pw": "new-open", "owner_pw": "new-owner", "fields": {"金额": "100"}}])
    _pick_event_stub(app, ev["id"])
    monkeypatch.setattr(app, "_pool_tag_dialog", lambda *a, **kw: None)
    app.open_event(ev["id"])
    getattr(app, entry)()
    saved = app.store.list_pool()[0]
    assert saved["cc"] == "cc@example.com" and saved["fields"] == {}
    assert saved["open_pw"] == "old-open" and saved["owner_pw"] == "old-owner"
    assert app.store.load_event(ev["id"])[0]["recipients"][0]["fields"] == {"金额": "100"}


def test_pool_import_save_failure_does_not_report_success(app, tmp_path, monkeypatch, no_messagebox):
    _pool_rec(app, "甲")
    path = tmp_path / "failure.xlsx"
    _file(path, [["收件人名称", "收件人", "抄送"], ["甲", "a@example.com", "cc@example.com"]])
    monkeypatch.setattr(app.store, "save_pool_recipient", lambda rec: (False, "test write failure", rec))
    _import(app, monkeypatch, path)
    assert ("error", "保存失败") in no_messagebox
    assert ("info", "导入完成") not in no_messagebox
    assert app.store.list_pool()[0]["cc"] == ""


def test_pool_template_includes_tags_but_not_business_fields(app, tmp_path, monkeypatch):
    path = tmp_path / "template.xlsx"
    monkeypatch.setattr(lp_tool.filedialog, "asksaveasfilename", lambda **kw: str(path))
    _buttons(app.home_tab_pool)["下载联系人模板"].invoke()
    headers = _read(path)[0]
    assert "标签" in headers
    assert not {"Call_Amount", "Unfunded", "Due_Date"} & set(headers)
    assert "标签" not in lp_v3.import_template_headers(["Call_Amount"])
